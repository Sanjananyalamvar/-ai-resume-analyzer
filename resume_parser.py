import io
import os
import re
import shutil

MIN_TEXT_CHARS = 80      # fewer letters than this means "no readable text"
OCR_MAX_PAGES = 3        # resumes are short, so only the first pages are scanned


class NeedsOCRError(ValueError):
    """The file has no selectable text and OCR (text recognition) is not available."""


# ---------------------------------------------------------------- helpers
def _read_bytes(file):
    """Return the file content as bytes (works for Streamlit uploads and open files)."""
    if hasattr(file, "getvalue"):
        return file.getvalue()
    file.seek(0)
    return file.read()


def _enough(text):
    return len(re.sub(r"\s", "", text or "")) >= MIN_TEXT_CHARS


# ------------------------------------------------ two-column aware PDF reading
def _group_lines(words, tol=3.0):
    """Group words into lines by their vertical position."""
    lines, current, ref = [], [], None
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if ref is None or abs(w["top"] - ref) <= tol:
            current.append(w)
            ref = w["top"] if ref is None else ref
        else:
            lines.append(sorted(current, key=lambda w: w["x0"]))
            current, ref = [w], w["top"]
    if current:
        lines.append(sorted(current, key=lambda w: w["x0"]))
    return lines


def _line_text(line):
    return " ".join(w["text"] for w in line)


def _find_gutter(words, page_width):
    """x position of the empty strip between two columns, or None for one column."""
    n = len(words)
    if n < 30:
        return None
    allowed = max(1, int(0.05 * n))          # a few full-width header words may cross
    lo, hi = int(0.28 * page_width), int(0.72 * page_width)

    best_run, run = (0, None), []
    for x in range(lo, hi + 1):
        crossing = sum(1 for w in words if w["x0"] < x < w["x1"])
        if crossing <= allowed:
            run.append(x)
        else:
            if len(run) > best_run[0]:
                best_run = (len(run), run[len(run) // 2])
            run = []
    if len(run) > best_run[0]:
        best_run = (len(run), run[len(run) // 2])

    width, split = best_run
    if split is None or width < 12:           # a real gutter is at least ~12 points wide
        return None
    left = sum(1 for w in words if (w["x0"] + w["x1"]) / 2 < split)
    if left < 0.15 * n or (n - left) < 0.15 * n:
        return None
    return split


def order_words(words, page_width):
    """Return page text in reading order, reading each column top to bottom."""
    if not words:
        return ""
    lines = _group_lines(words)
    split = _find_gutter(words, page_width)
    if split is None:
        return "\n".join(_line_text(l) for l in lines)

    out, region = [], []

    def flush():
        if not region:
            return
        flat = [w for l in region for w in l]
        for column in ([w for w in flat if (w["x0"] + w["x1"]) / 2 < split],
                       [w for w in flat if (w["x0"] + w["x1"]) / 2 >= split]):
            out.extend(_line_text(l) for l in _group_lines(column))
        region.clear()

    for line in lines:
        if any(w["x0"] < split < w["x1"] for w in line):   # full-width line (e.g. the header)
            flush()
            out.append(_line_text(line))
        else:
            region.append(line)
    flush()
    return "\n".join(out)


def _pdf_text_layout(data):
    import pdfplumber
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(order_words(page.extract_words(), float(page.width)))
    return "\n".join(pages)


def _pdf_text_plain(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


# ------------------------------------------------------------------- OCR
def ocr_available():
    """True if Tesseract and its Python wrapper can be used."""
    try:
        import pytesseract
        from PIL import Image  # noqa: F401
    except Exception:
        return False
    cmd = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
    if not cmd:
        for path in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                     r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
            if os.path.exists(path):
                cmd = path
                break
    if not cmd:
        return False
    pytesseract.pytesseract.tesseract_cmd = cmd
    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_image(img):
    import pytesseract
    img = img.convert("RGB")
    longest = max(img.size)
    if longest > 3500:                       # very large photos are slow, so shrink them
        scale = 3500 / longest
        img = img.resize((int(img.width * scale), int(img.height * scale)))
    return pytesseract.image_to_string(img, timeout=90)


def _ocr_pdf(data):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(data)
    try:
        texts = []
        for i in range(min(len(pdf), OCR_MAX_PAGES)):
            page = pdf[i]
            texts.append(_ocr_image(page.render(scale=2.5).to_pil()))
        return "\n".join(texts)
    finally:
        pdf.close()


NO_OCR_MESSAGE = (
    "No selectable text was found. The file may be a scan, a photo or an image-only PDF "
    "(or it may be damaged), and text recognition (OCR) is not available here. "
    "Please upload a text-based PDF (export it from Word or Google Docs) or a DOCX file."
)


def _extract_pdf(data):
    layout_text = plain_text = ""
    errors = 0
    try:
        layout_text = _pdf_text_layout(data)
    except Exception:
        errors += 1
    try:
        plain_text = _pdf_text_plain(data)
    except Exception:
        errors += 1

    if _enough(layout_text):
        return layout_text
    if _enough(plain_text):
        return plain_text

    if ocr_available():
        try:
            ocr_text = _ocr_pdf(data)
        except Exception:
            raise ValueError("This PDF looks damaged and could not be read. "
                             "Please export it again and upload the new file.")
        if _enough(ocr_text):
            return ocr_text
        raise ValueError("Text recognition could not find readable text in this file. "
                         "Try a clearer scan or a text-based PDF.")

    if errors == 2:
        raise ValueError("This PDF looks damaged and could not be read. "
                         "Please export it again and upload the new file.")
    raise NeedsOCRError(NO_OCR_MESSAGE)


# ------------------------------------------------------------------ DOCX
def _docx_text(data):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    lines = []

    def add_paragraph(p):
        text = p.text.strip()
        if text:
            lines.append(text)

    def add_table(table):
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:         # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                for p in cell.paragraphs:
                    add_paragraph(p)
                for inner in cell.tables:
                    add_table(inner)

    for child in doc.element.body.iterchildren():   # keeps paragraphs and tables in order
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            add_paragraph(Paragraph(child, doc))
        elif tag == "tbl":
            add_table(Table(child, doc))

    # Text boxes, which fancy templates use for sidebars
    known = set(lines)
    for box in doc.element.body.iter(qn("w:txbxContent")):
        for p in box.iter(qn("w:p")):
            text = "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()
            if text and text not in known:
                known.add(text)
                lines.append(text)
    return "\n".join(lines)


# ----------------------------------------------------------------- images
def _extract_image(data):
    if not ocr_available():
        raise NeedsOCRError(NO_OCR_MESSAGE)
    from PIL import Image
    text = _ocr_image(Image.open(io.BytesIO(data)))
    if not _enough(text):
        raise ValueError("Text recognition could not find readable text in this image. "
                         "Try a sharper, straight-on photo.")
    return text


# ------------------------------------------------------------- main entry
def extract_text(file, filename):
    """Extract text from a PDF, DOCX, PNG or JPG resume."""
    name = filename.lower()
    data = _read_bytes(file)

    if name.endswith(".pdf"):
        text = _extract_pdf(data)
    elif name.endswith(".docx"):
        text = _docx_text(data)
    elif name.endswith((".png", ".jpg", ".jpeg")):
        text = _extract_image(data)
    else:
        raise ValueError("Only PDF, DOCX, PNG and JPG files are supported")

    if not _enough(text):
        raise ValueError("Very little text was found in this file. "
                         "It may be empty or an image. Try a text-based PDF or a DOCX file.")
    return text
